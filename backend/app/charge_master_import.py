"""Charge master import from an Excel (.xlsx) file.

Sheet layout (first row is the header, order does not matter, names are case-insensitive):

    category | code | name | price | active
    consultation | C001 | Consultation, general | 30.00 | yes
    misc | M010 | Medical report | 50 |

* category: `consultation` or `misc` (required)
* name: required, up to 255 characters
* price: RM, a number with at most 2 decimals, 0 to 1,000,000 (required)
* code: optional, up to 64 characters
* active: optional, yes/no (default yes)

Import only ever APPENDS. A row whose category + code (or, when either side has no code,
category + name) already exists is skipped and counted, never overwritten. The file is all or
nothing: any bad row means nothing is added, and the errors say which rows.
"""
import io
import zipfile
from decimal import Decimal, InvalidOperation
from typing import List, Optional, Tuple

from openpyxl import Workbook, load_workbook

MAX_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED = 10 * 1024 * 1024
MAX_SCANNED_ROWS = 50_000  # blank formatted rows must not run forever
MAX_ROWS = 2000
MAX_ERRORS = 50
HEADERS = ("category", "code", "name", "price", "active")
_REQUIRED = ("category", "name", "price")
_MAX_PRICE_CENTS = 100_000_000
_YES = {"yes", "y", "true", "1", "active"}
_NO = {"no", "n", "false", "0", "inactive"}


class ImportFileError(ValueError):
    """The file as a whole cannot be read (wrong type, no header, too big)."""


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Charge master"
    ws.append(list(HEADERS))
    ws.append(["consultation", "C001", "Consultation, general", 30.00, "yes"])
    ws.append(["misc", "M010", "Medical report", 50, "yes"])
    for col, width in zip("ABCDE", (16, 12, 40, 10, 10)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_text(v) -> str:
    return "" if v is None else str(v).strip()


def _price_cents(v) -> int:
    if isinstance(v, bool):
        raise ValueError("price must be a number")
    if isinstance(v, (int, float, Decimal)):
        d = Decimal(str(v))
    else:
        t = _cell_text(v).replace(",", "")
        if t.upper().startswith("RM"):
            t = t[2:].strip()
        try:
            d = Decimal(t)
        except InvalidOperation:
            raise ValueError("price must be a number, for example 30 or 30.50")
    if not d.is_finite():
        raise ValueError("price must be a number")
    if d > _MAX_PRICE_CENTS // 100:  # before multiplying: 1e999999 would overflow Decimal
        raise ValueError("price is above RM 1,000,000")
    if d < 0:
        raise ValueError("price cannot be negative")
    cents = d * 100
    if cents != cents.to_integral_value():
        raise ValueError("price can have at most 2 decimals")
    cents = int(cents)
    if cents > _MAX_PRICE_CENTS:
        raise ValueError("price is above RM 1,000,000")
    return cents


def parse_workbook(raw: bytes, categories) -> Tuple[List[dict], List[dict]]:
    """Returns (rows, errors). rows: {row, category, code, name, unit_price_cents, active}.
    errors: {row, message} (row 1 is the header, so data starts at 2)."""
    if len(raw) > MAX_BYTES:
        raise ImportFileError(f"File too large, max {MAX_BYTES // (1024 * 1024)} MB.")
    if not raw.startswith(b"PK"):
        raise ImportFileError("This is not an .xlsx file. Save the sheet as Excel Workbook (.xlsx).")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise ImportFileError("The workbook is too large once unpacked.")
    except zipfile.BadZipFile:
        raise ImportFileError("Could not open the file as an Excel workbook.")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Could not open the file as an Excel workbook.")
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            raise ImportFileError("The first sheet is empty.")
        cols = {_cell_text(h).lower(): i for i, h in enumerate(header) if _cell_text(h)}
        missing = [h for h in _REQUIRED if h not in cols]
        if missing:
            raise ImportFileError(
                "Missing column(s) in the first row: " + ", ".join(missing)
                + ". The first row must contain: " + ", ".join(HEADERS) + ".")

        rows, errors = [], []
        used = 0
        for n, r in enumerate(it, start=2):
            if n > MAX_SCANNED_ROWS:
                raise ImportFileError("The sheet has too many rows, remove the empty formatted rows.")

            def get(name):
                i = cols.get(name)
                return r[i] if i is not None and i < len(r) else None

            if all(_cell_text(c) == "" for c in r):
                continue  # blank line
            used += 1
            if used > MAX_ROWS:
                raise ImportFileError(f"Too many rows, max {MAX_ROWS}.")
            problems = []
            category = _cell_text(get("category")).lower()
            if category not in categories:
                problems.append(f"category must be one of: {', '.join(sorted(categories))}")
            name = _cell_text(get("name"))
            if not name:
                problems.append("name is required")
            elif len(name) > 255:
                problems.append("name is longer than 255 characters")
            code = _cell_text(get("code")) or None
            if code and len(code) > 64:
                problems.append("code is longer than 64 characters")
            cents = 0
            try:
                cents = _price_cents(get("price"))
            except ValueError as e:
                problems.append(str(e))
            active_raw = _cell_text(get("active")).lower()
            active = True
            if active_raw:
                if active_raw in _YES:
                    active = True
                elif active_raw in _NO:
                    active = False
                else:
                    problems.append("active must be yes or no")
            if problems:
                errors.append({"row": n, "message": "; ".join(problems)})
            else:
                rows.append({"row": n, "category": category, "code": code, "name": name,
                             "unit_price_cents": cents, "active": active})
        if not rows and not errors:
            raise ImportFileError("There are no data rows under the header.")
        return rows, errors
    finally:
        wb.close()


def same_item(a: dict, b: dict) -> bool:
    """Duplicate rule: same category and same code, or (when either has no code) same name."""
    if a["category"] != b["category"]:
        return False
    if a.get("code") and b.get("code"):
        return a["code"].lower() == b["code"].lower()
    return a["name"].lower() == b["name"].lower()


def split_new_and_duplicates(rows: List[dict], existing: List[dict]) -> Tuple[List[dict], List[dict]]:
    """Rows to add, and rows skipped as duplicates of an existing item or an earlier row."""
    seen = list(existing)
    new, dups = [], []
    for row in rows:
        if any(same_item(row, s) for s in seen):
            dups.append(row)
        else:
            new.append(row)
            seen.append(row)
    return new, dups


def first_errors(errors: List[dict]) -> Optional[List[dict]]:
    return errors[:MAX_ERRORS] if errors else None
